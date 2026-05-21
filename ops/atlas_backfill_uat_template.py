#!/usr/bin/env python3
"""
Backfill Atlas UAT template columns using queryOrderDetails.do.

This script updates only cell values (keeps the template format intact):
- "Transaction Fee (Currency + Amount)"
- "Airline PNR"

Prereqs (recommended via env vars):
  ATLAS_BASE_URL=https://sandbox.atriptech.com
  ATLAS_CLIENT_ID=...
  ATLAS_CLIENT_SECRET=...
  ATLAS_REQUEST_SOURCE=TourValve   (optional)
"""

from __future__ import annotations

import argparse
import json
import os
import time
import urllib.error
import urllib.request
from typing import Any

from openpyxl import load_workbook


def _atlas_post(
    base_url: str,
    path: str,
    client_id: str,
    client_secret: str,
    payload: dict[str, Any],
    timeout_s: int,
) -> dict[str, Any]:
    url = base_url.rstrip("/") + path
    req = urllib.request.Request(
        url,
        data=json.dumps(payload).encode("utf-8"),
        method="POST",
    )
    req.add_header("Content-Type", "application/json")
    req.add_header("Accept", "application/json")
    req.add_header("x-atlas-client-id", client_id)
    req.add_header("x-atlas-client-secret", client_secret)

    with urllib.request.urlopen(req, timeout=timeout_s) as resp:
        body = resp.read().decode("utf-8", errors="replace")

    data = json.loads(body)
    if isinstance(data, dict) and data.get("status") not in (0, None):
        raise RuntimeError(f"Atlas API error {data.get('status')}: {data.get('msg')}")
    return data


def _extract_airline_pnrs(details: dict[str, Any]) -> list[str]:
    pax_ticket_infos = details.get("paxTicketInfos") or []
    airline_pnrs: list[str] = []
    if isinstance(pax_ticket_infos, list):
        for item in pax_ticket_infos:
            if not isinstance(item, dict):
                continue
            pnrs = item.get("airlinePNRs") or []
            if isinstance(pnrs, list):
                airline_pnrs.extend([str(x).strip() for x in pnrs if str(x).strip()])

    legacy = details.get("airPnr")
    if legacy:
        for piece in str(legacy).replace(";", ",").split(","):
            piece = piece.strip()
            if piece:
                airline_pnrs.append(piece)

    # stable order, no duplicates
    seen: set[str] = set()
    out: list[str] = []
    for pnr in airline_pnrs:
        if pnr in seen:
            continue
        seen.add(pnr)
        out.append(pnr)
    return out


def _format_fee(currency: str | None, amount: Any) -> str | None:
    if amount is None:
        return None
    try:
        value = float(amount)
    except Exception:
        return None

    currency_code = (currency or "").strip() or "MYR"
    return f"{currency_code} {value:.2f}"


def _find_col_idx(header_row: list[Any], header_name: str) -> int:
    normalized = header_name.strip().lower()
    for idx, cell in enumerate(header_row, start=1):
        if cell is None:
            continue
        if str(cell).strip().lower() == normalized:
            return idx
    raise KeyError(f"Header not found: {header_name}")


def main() -> int:
    parser = argparse.ArgumentParser(description="Backfill Atlas UAT template from queryOrderDetails.do")
    parser.add_argument("--in", dest="in_path", required=True, help="Input .xlsx path")
    parser.add_argument("--out", dest="out_path", required=True, help="Output .xlsx path")
    parser.add_argument("--base-url", default=os.getenv("ATLAS_BASE_URL", "").strip())
    parser.add_argument("--client-id", default=os.getenv("ATLAS_CLIENT_ID", "").strip())
    parser.add_argument("--client-secret", default=os.getenv("ATLAS_CLIENT_SECRET", "").strip())
    parser.add_argument("--request-source", default=os.getenv("ATLAS_REQUEST_SOURCE", "TourValve").strip())
    parser.add_argument("--timeout", type=int, default=30)
    parser.add_argument("--sleep-ms", type=int, default=300, help="Delay between Atlas calls")
    parser.add_argument("--poll-attempts", type=int, default=0, help="Extra queryOrderDetails retries until ticketed")
    parser.add_argument("--poll-delay-s", type=float, default=5.0, help="Delay between poll attempts")
    args = parser.parse_args()

    if not args.base_url:
        raise SystemExit("Missing --base-url (or ATLAS_BASE_URL).")
    if not args.client_id:
        raise SystemExit("Missing --client-id (or ATLAS_CLIENT_ID).")
    if not args.client_secret:
        raise SystemExit("Missing --client-secret (or ATLAS_CLIENT_SECRET).")

    wb = load_workbook(args.in_path)
    ws = wb[wb.sheetnames[0]]

    header = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    col_order_no = _find_col_idx(header, "OrderNo")
    col_airline_pnr = _find_col_idx(header, "Airline PNR")
    col_tx_fee = _find_col_idx(header, "Transaction Fee (Currency + Amount)")

    updated = 0
    for row in range(2, ws.max_row + 1):
        order_no = ws.cell(row, col_order_no).value
        if not order_no:
            continue
        order_no = str(order_no).strip()
        if not order_no:
            continue

        details: dict[str, Any] | None = None
        for poll_attempt in range(max(args.poll_attempts, 0) + 1):
            try:
                details = _atlas_post(
                    args.base_url,
                    "/queryOrderDetails.do",
                    args.client_id,
                    args.client_secret,
                    {"orderNo": order_no, "requestSource": args.request_source},
                    timeout_s=args.timeout,
                )
            except (urllib.error.URLError, urllib.error.HTTPError, RuntimeError, ValueError) as ex:
                # Keep template intact; just skip problematic rows.
                print(f"[WARN] {order_no}: {ex}")
                details = None
                break

            # Stop polling once ticketed.
            raw_status = details.get("orderStatus")
            try:
                status_int = int(raw_status) if raw_status is not None else None
            except Exception:
                status_int = None

            airline_pnrs_now = _extract_airline_pnrs(details)
            if status_int == 2 or airline_pnrs_now:
                break

            if poll_attempt < args.poll_attempts:
                time.sleep(max(args.poll_delay_s, 0.1))

        if details is None:
            continue

        currency = details.get("currency") or (details.get("routing") or {}).get("currency")
        tx_fee = _format_fee(currency, details.get("totalTransactionFee"))
        airline_pnrs = _extract_airline_pnrs(details)

        if tx_fee:
            ws.cell(row, col_tx_fee).value = tx_fee
        if airline_pnrs:
            ws.cell(row, col_airline_pnr).value = ", ".join(airline_pnrs)

        updated += 1
        if args.sleep_ms > 0:
            time.sleep(args.sleep_ms / 1000.0)

    wb.save(args.out_path)
    print(f"Updated rows: {updated}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
